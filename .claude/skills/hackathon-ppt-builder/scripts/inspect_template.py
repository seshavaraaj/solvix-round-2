#!/usr/bin/env python3
"""Inspect a hackathon template deck before filling it.

Usage:
    python inspect_template.py template.pptx [--json out.json]

Prints, per slide:
  - role guess (title / content / instructions / thank-you)
  - every text-bearing shape: name, position/size (inches), font sizes used,
    current text, and a rough capacity estimate (how many characters of body
    text fit at the shape's font size)
  - pictures, tables, charts, groups (so you know what visuals already exist)
Then prints every sentence in the deck (slides + notes) that looks like a
slide/page limit, time limit, or format rule (e.g. "maximum 6 slides",
"submit as PDF", "5 minutes to present").

Capacity is an estimate (average glyph width ~0.5 em, line height 1.2 em).
Treat it as an upper bound and aim for ~70% of it.
"""
import argparse
import json
import re
import sys

from pptx import Presentation
from pptx.util import Emu

EMU_PER_IN = 914400

LIMIT_PATTERNS = [
    r"(max(imum)?|up\s*to|not\s+(to\s+)?exceed|no\s+more\s+than|limit(ed)?\s+(of|to)?|at\s+most)[^.\n]{0,40}\b(\d+)\b[^.\n]{0,20}(slides?|pages?)",
    r"\b(\d+)\b\s*(slides?|pages?)\s*(max(imum)?|only|limit)",
    r"(slides?|pages?)\s*limit[^.\n]{0,30}\b(\d+)\b",
    r"\b(\d+)\s*(min(ute)?s?)\b[^.\n]{0,40}(present|pitch|demo)",
    r"(present|pitch|demo)[^.\n]{0,40}\b(\d+)\s*(min(ute)?s?)\b",
    r"\b(pdf|ppt|pptx)\b[^.\n]{0,40}(format|submit|save|only)",
    r"(submit|save)[^.\n]{0,40}\b(pdf|ppt|pptx)\b",
    r"(avoid\s+paragraphs|in\s+points|bullet|infographic|diagram|flow\s*chart)",
    r"(do\s+not|don'?t|without)\s+(change|chang(ing)?|modify|alter)[^.\n]{0,60}",
    r"\b(\d+)\s+of\s+(\d+)\b",
]
INSTRUCTION_HINTS = re.compile(
    r"(important\s+instructions|guidelines|please\s+ensure|kindly|note\s*[-–:]|instructions)",
    re.I,
)


def inches(v):
    return round(Emu(v).inches, 2) if v is not None else None


def font_sizes(shape):
    sizes = set()
    if not shape.has_text_frame:
        return []
    for p in shape.text_frame.paragraphs:
        for r in p.runs:
            if r.font.size:
                sizes.add(r.font.size.pt)
    return sorted(sizes)


def capacity(shape, pt):
    """Rough character capacity of a text shape at a given font size."""
    if shape.width is None or shape.height is None or not pt:
        return None
    w_in = shape.width / EMU_PER_IN - 0.2  # default insets
    h_in = shape.height / EMU_PER_IN - 0.1
    if w_in <= 0 or h_in <= 0:
        return None
    chars_per_line = (w_in * 72) / (pt * 0.5)
    lines = (h_in * 72) / (pt * 1.2)
    return int(chars_per_line * lines), int(chars_per_line), int(lines)


def walk(shapes, depth=0):
    for s in shapes:
        yield s, depth
        if s.shape_type == 6:  # group
            yield from walk(s.shapes, depth + 1)


def guess_role(idx, total, texts):
    joined = " ".join(texts).lower()
    if INSTRUCTION_HINTS.search(joined) and ("slide" in joined or "ppt" in joined or "submit" in joined):
        return "instructions (usually removed before submission)"
    if idx == 1:
        return "title"
    if idx == total and re.search(r"thank|questions|q\s*&\s*a", joined):
        return "closing"
    return "content"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("deck")
    ap.add_argument("--json", help="also write the inventory as JSON here")
    a = ap.parse_args()

    prs = Presentation(a.deck)
    sw, sh = inches(prs.slide_width), inches(prs.slide_height)
    total = len(prs.slides)
    print(f"Deck: {a.deck}\nSlide size: {sw} x {sh} in   Slides: {total}\n")

    inventory = {"slide_width_in": sw, "slide_height_in": sh, "slides": [], "rules": []}
    all_text = []

    for i, slide in enumerate(prs.slides, 1):
        texts, shapes_out = [], []
        for s, depth in walk(slide.shapes):
            kind = str(s.shape_type).split(".")[-1].split(" ")[0] if s.shape_type else "SHAPE"
            entry = {
                "name": s.name,
                "kind": kind,
                "depth": depth,
                "is_placeholder": s.is_placeholder,
                "x": inches(s.left), "y": inches(s.top),
                "w": inches(s.width), "h": inches(s.height),
            }
            if s.has_text_frame and s.text_frame.text.strip():
                t = s.text_frame.text.strip()
                texts.append(t)
                sizes = font_sizes(s)
                entry["text"] = t
                entry["font_sizes_pt"] = sizes
                entry["paragraphs"] = len(s.text_frame.paragraphs)
                body_pt = min(sizes) if sizes else 18.0
                cap = capacity(s, body_pt)
                if cap:
                    entry["capacity_chars"], entry["chars_per_line"], entry["lines"] = cap
                entry["current_chars"] = len(t)
            if getattr(s, "has_table", False) and s.has_table:
                entry["table"] = f"{len(s.table.rows)}x{len(s.table.columns)}"
            if getattr(s, "has_chart", False) and s.has_chart:
                entry["chart"] = True
            shapes_out.append(entry)

        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        role = guess_role(i, total, texts)
        inventory["slides"].append(
            {"index": i, "file": str(slide.part.partname).lstrip("/"), "layout": slide.slide_layout.name,
             "role": role, "shapes": shapes_out, "notes": notes}
        )
        all_text += [(i, t) for t in texts] + ([(i, notes)] if notes else [])

        print(f"=== Slide {i}  [{role}]  {str(slide.part.partname).lstrip('/')}  layout='{slide.slide_layout.name}'")
        for e in shapes_out:
            pad = "  " * (e["depth"] + 1)
            geo = f"@({e['x']},{e['y']}) {e['w']}x{e['h']}in"
            if "text" in e:
                cap = e.get("capacity_chars")
                capstr = f" cap~{cap}ch ({e.get('chars_per_line')}/line x {e.get('lines')} lines)" if cap else ""
                preview = e["text"].replace("\n", " | ")
                if len(preview) > 160:
                    preview = preview[:157] + "..."
                print(f"{pad}- {e['name']} [{e['kind']}{', ph' if e['is_placeholder'] else ''}] {geo} "
                      f"fonts={e['font_sizes_pt'] or 'inherited'}{capstr}")
                print(f"{pad}  \"{preview}\"")
            else:
                extra = e.get("table") or ("chart" if e.get("chart") else "")
                print(f"{pad}- {e['name']} [{e['kind']}] {geo} {extra}")
        if notes:
            print(f"  notes: {notes[:200]}")
        print()

    print("=== Rules / limits found in the deck text (verify each by reading it in context)")
    seen = set()
    for idx, t in all_text:
        for sentence in re.split(r"(?<=[.!?])\s+|\n", t):
            s_clean = sentence.strip()
            if not s_clean or s_clean in seen:
                continue
            for pat in LIMIT_PATTERNS:
                if re.search(pat, s_clean, re.I):
                    seen.add(s_clean)
                    inventory["rules"].append({"slide": idx, "text": s_clean})
                    print(f"  slide {idx}: {s_clean}")
                    break
    if not inventory["rules"]:
        print("  (none found — if a slide increase is needed, ask the user for the limit)")

    if a.json:
        with open(a.json, "w") as f:
            json.dump(inventory, f, indent=2)
        print(f"\nJSON inventory written to {a.json}")


if __name__ == "__main__":
    sys.exit(main())
